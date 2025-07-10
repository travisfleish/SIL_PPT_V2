# scripts/debug_template_slides.py
"""
Debug script to inspect the PowerPoint template and find where static slides are located
"""

from pptx import Presentation
from pathlib import Path
import logging

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


def debug_template_structure():
    """
    Inspect the template file to understand its structure
    """
    template_path = Path('templates/sil_combined_template.pptx')

    if not template_path.exists():
        print(f"❌ Template not found at: {template_path}")
        return

    try:
        presentation = Presentation(str(template_path))

        print("🔍 TEMPLATE STRUCTURE ANALYSIS")
        print("=" * 50)

        # 1. Analyze slide layouts
        print(f"\n📐 SLIDE LAYOUTS ({len(presentation.slide_layouts)} total):")
        for i, layout in enumerate(presentation.slide_layouts):
            print(f"  Layout {i}: {layout.name}")

        # 2. Analyze actual slides
        print(f"\n📄 ACTUAL SLIDES ({len(presentation.slides)} total):")
        for i, slide in enumerate(presentation.slides):
            slide_layout_name = slide.slide_layout.name if slide.slide_layout else "Unknown"

            # Try to get slide title or first text
            slide_text = "No text found"
            try:
                for shape in slide.shapes:
                    if hasattr(shape, 'text_frame') and shape.has_text_frame:
                        if shape.text_frame.text.strip():
                            slide_text = shape.text_frame.text.strip()[:50] + "..."
                            break
            except:
                pass

            print(f"  Slide {i}: Layout='{slide_layout_name}' | Text='{slide_text}'")

        # 3. Look for slides that might be our static slides
        print(f"\n🎯 POTENTIAL STATIC SLIDES:")
        static_slide_keywords = [
            "how to use", "sports innovation lab", "sil", "branding",
            "report", "insights", "playbook", "sponsorship"
        ]

        for i, slide in enumerate(presentation.slides):
            slide_text_combined = ""
            try:
                for shape in slide.shapes:
                    if hasattr(shape, 'text_frame') and shape.has_text_frame:
                        slide_text_combined += shape.text_frame.text.lower() + " "
            except:
                pass

            # Check if this slide contains static content keywords
            for keyword in static_slide_keywords:
                if keyword in slide_text_combined:
                    print(f"  🎯 Slide {i} might be static - contains '{keyword}'")
                    print(f"      Text preview: {slide_text_combined[:100]}...")
                    break

        # 4. Generate the corrected code
        print(f"\n🔧 RECOMMENDED FIX:")
        if len(presentation.slides) >= 2:
            print("Based on your template structure, update your pptx_builder.py:")
            print("\n# In _copy_static_slide_from_template method, change:")
            print("# Old:")
            print("#   template_slide = template_presentation.slides[template_slide_index]")
            print("# New:")
            if len(presentation.slides) > 13:
                print(
                    f"#   Static slide positions appear to be: {len(presentation.slides) - 2} and {len(presentation.slides) - 1}")
            else:
                print(
                    f"#   Static slide positions appear to be: 0 and 1 (or {len(presentation.slides) - 2} and {len(presentation.slides) - 1})")

        print(f"\n✅ Analysis complete! Check the positions above.")

    except Exception as e:
        print(f"❌ Error analyzing template: {e}")


def suggest_fix(how_to_slide_pos: int, branding_slide_pos: int):
    """
    Generate the corrected code with the right slide positions
    """
    print(f"\n🔧 CORRECTED CODE:")
    print("Replace the build_presentation method in pptx_builder.py with:")
    print(f"""
# 2. Add "How To Use This Report" static slide (from template slide {how_to_slide_pos})
self._copy_static_slide_from_template({how_to_slide_pos}, "How To Use This Report")

# ... (rest of your slides)

# 7. Add SIL branding slide at the end (from template slide {branding_slide_pos})
self._copy_static_slide_from_template({branding_slide_pos}, "Sports Innovation Lab Branding")
""")


if __name__ == "__main__":
    debug_template_structure()

    print(f"\n📋 MANUAL VERIFICATION STEPS:")
    print("1. Open your template file: templates/sil_combined_template.pptx")
    print("2. Look at the slide thumbnails on the left")
    print("3. Identify which slide numbers contain:")
    print("   - 'How To Use This Report' content")
    print("   - 'Sports Innovation Lab' branding")
    print("4. Note that PowerPoint counts slides starting from 1, but Python starts from 0")
    print("5. Update the slide positions in the debug output above")

    print(f"\n💡 QUICK TEST:")
    print("Run this script to see your template structure, then update the positions accordingly.")